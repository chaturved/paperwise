import io
import shutil
import zipfile

import docx
import fitz
import pytest

from src.config import settings
from src.core.exceptions import BadRequestError
from src.integrations.documents import Page
from src.integrations.documents.ocr import OcrUnavailableError, PageOcr
from src.integrations.documents.parsers import DocumentParser, PdfParser, TextBudget
from src.integrations.documents.timeouts import ExtractionTimeout
from src.utils.text_utils import TextChunker


def _docx_bytes(build) -> bytes:
    document = docx.Document()
    build(document)
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


class TestExtractDocx:
    def test_keeps_paragraphs_and_table_cells_in_document_order(self):
        def build(document):
            document.add_paragraph("Payment terms")
            table = document.add_table(rows=2, cols=2)
            table.cell(0, 0).text = "Due"
            table.cell(0, 1).text = "30 days"
            table.cell(1, 0).text = "Late fee"
            table.cell(1, 1).text = "2%"
            document.add_paragraph("Termination")

        [page] = DocumentParser().parse(_docx_bytes(build), "docx")

        assert page.number is None
        assert page.text.splitlines() == ["Payment terms", "Due | 30 days", "Late fee | 2%", "Termination"]

    def test_empty_document_has_no_readable_text(self):
        assert [TextChunker().clean(p.text) for p in DocumentParser().parse(_docx_bytes(lambda d: None), "docx")] == [
            ""
        ]


def _zip_bytes(members: dict[str, bytes]) -> bytes:
    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w", zipfile.ZIP_DEFLATED) as archive:
        for name, data in members.items():
            archive.writestr(name, data)
    return buffer.getvalue()


class TestUploadSafetyLimits:
    def test_docx_that_unpacks_far_beyond_its_size_is_rejected(self):
        # Zeros compress to almost nothing, so the file is tiny but declares a huge unpacked size.
        bomb = _zip_bytes({"word/document.xml": b"\0" * (settings.MAX_DOCX_UNZIPPED_MB * 1024 * 1024 + 1)})

        assert len(bomb) < 1024 * 1024
        with pytest.raises(BadRequestError) as exc:
            DocumentParser().parse(bomb, "docx")
        assert "too large" in exc.value.message

    def test_docx_with_too_many_entries_is_rejected(self):
        many = _zip_bytes({f"part{i}.xml": b"x" for i in range(settings.MAX_DOCX_ENTRIES + 1)})

        with pytest.raises(BadRequestError) as exc:
            DocumentParser().parse(many, "docx")
        assert "too large" in exc.value.message

    def test_pdf_with_too_many_pages_is_rejected(self, monkeypatch):
        monkeypatch.setattr(settings, "MAX_PDF_PAGES", 2)
        pdf = fitz.open()
        for _ in range(3):
            pdf.new_page()
        contents = pdf.tobytes()
        pdf.close()

        with pytest.raises(BadRequestError) as exc:
            DocumentParser().parse(contents, "pdf")
        assert "too large" in exc.value.message

    def test_forged_size_in_zip_headers_is_not_trusted(self):
        # Declare a tiny unpacked size in both the local and central headers while the data stays large.
        real = b"A" * 5000
        buffer = io.BytesIO()
        with zipfile.ZipFile(buffer, "w", zipfile.ZIP_STORED) as archive:
            archive.writestr("word/document.xml", real)
        forged = bytearray(buffer.getvalue())
        local = forged.index(b"PK\x03\x04")
        central = forged.index(b"PK\x01\x02")
        forged[local + 22 : local + 26] = (10).to_bytes(4, "little")
        forged[central + 24 : central + 28] = (10).to_bytes(4, "little")

        with pytest.raises(BadRequestError):
            DocumentParser().parse(bytes(forged), "docx")

    def test_decompressed_output_over_the_cap_stops_early(self, monkeypatch):
        monkeypatch.setattr(settings, "MAX_DOCX_UNZIPPED_MB", 1)
        two_parts = _zip_bytes({"a.xml": b"\0" * (700 * 1024), "b.xml": b"\0" * (700 * 1024)})

        with pytest.raises(BadRequestError) as exc:
            DocumentParser().parse(two_parts, "docx")
        assert "too large" in exc.value.message

    def test_pdf_text_over_the_cap_is_rejected_while_reading_pages(self, monkeypatch):
        monkeypatch.setattr(settings, "MAX_EXTRACTED_CHARS", 50)
        pdf = fitz.open()
        for _ in range(5):
            pdf.new_page().insert_text((72, 72), "word " * 10)
        contents = pdf.tobytes()
        pdf.close()

        with pytest.raises(BadRequestError) as exc:
            DocumentParser().parse(contents, "pdf")
        assert "too large" in exc.value.message

    def test_text_longer_than_the_cap_is_rejected(self, monkeypatch):
        monkeypatch.setattr(settings, "MAX_EXTRACTED_CHARS", 10)

        with pytest.raises(BadRequestError) as exc:
            DocumentParser().parse(b"a" * 11, "txt")
        assert "too large" in exc.value.message

    def test_text_at_the_cap_is_accepted(self, monkeypatch):
        monkeypatch.setattr(settings, "MAX_EXTRACTED_CHARS", 10)

        assert DocumentParser().parse(b"a" * 10, "txt") == [Page(None, "a" * 10)]


class TestExtractOtherTypes:
    def test_txt_is_decoded(self):
        assert DocumentParser().parse(b"hello", "txt") == [Page(None, "hello")]

    def test_pdf_pages_are_numbered_from_one(self):
        pdf = fitz.open()
        for text in ("first page", "second page"):
            pdf.new_page().insert_text((72, 72), text)

        pages = DocumentParser().parse(pdf.tobytes(), "pdf")

        assert [(p.number, p.text.strip()) for p in pages] == [(1, "first page"), (2, "second page")]

    def test_unsupported_type_raises(self):
        with pytest.raises(BadRequestError):
            DocumentParser().parse(b"x", "exe")

    @pytest.mark.parametrize("ext", ["pdf", "docx"])
    def test_empty_or_corrupt_file_gives_a_plain_message(self, ext):
        for contents in (b"", b"not a real document"):
            with pytest.raises(BadRequestError) as exc:
                DocumentParser().parse(contents, ext)
            assert "couldn't be read" in exc.value.message


def _pdf(*pages: tuple[str, str]) -> bytes:
    """A PDF made of ("text", words) pages that have a text layer and ("image", words) pages that are only a picture."""
    target = fitz.open()
    for kind, words in pages:
        if kind == "text":
            target.new_page().insert_text((72, 100), words, fontsize=14)
            continue
        source = fitz.open()
        page = source.new_page()
        if words:
            page.insert_text((72, 120), words, fontsize=28)
        image_page = target.new_page()
        image_page.insert_image(image_page.rect, pixmap=page.get_pixmap(dpi=150))
    return target.tobytes()


class _FakeOcr(PageOcr):
    def __init__(self) -> None:
        self.read_pages: list[int] = []

    def read(self, page: fitz.Page) -> str:
        self.read_pages.append(page.number)
        return f"scanned page {page.number + 1}"


class _MissingOcr(PageOcr):
    def read(self, page: fitz.Page) -> str:
        raise OcrUnavailableError("Tesseract is not installed")


class TestScannedPdf:
    def test_pages_with_no_text_are_read_with_ocr_and_keep_their_page_numbers(self):
        ocr = _FakeOcr()
        contents = _pdf(("text", "Payment is due in 30 days"), ("image", "x"))

        pages = PdfParser(TextBudget(), ocr).parse(contents)

        assert [(p.number, p.text.strip()) for p in pages] == [(1, "Payment is due in 30 days"), (2, "scanned page 2")]
        assert ocr.read_pages == [1]

    def test_only_the_pages_read_with_ocr_are_marked_as_scanned(self):
        pages = PdfParser(TextBudget(), _FakeOcr()).parse(_pdf(("text", "Payment is due"), ("image", "x")))

        assert [p.scanned for p in pages] == [False, True]

    def test_ocr_is_not_used_when_every_page_has_text(self):
        ocr = _FakeOcr()

        PdfParser(TextBudget(), ocr).parse(_pdf(("text", "one"), ("text", "two")))

        assert ocr.read_pages == []

    def test_a_scan_with_too_many_pages_is_rejected_with_a_plain_message(self, monkeypatch):
        monkeypatch.setattr(settings, "OCR_MAX_PAGES", 2)
        contents = _pdf(("image", "a"), ("image", "b"), ("image", "c"))

        with pytest.raises(BadRequestError) as exc:
            PdfParser(TextBudget(), _FakeOcr()).parse(contents)

        assert "more than 2 pages" in exc.value.message

    def test_a_mostly_text_document_reads_blank_pages_only_up_to_the_cap(self, monkeypatch):
        monkeypatch.setattr(settings, "OCR_MAX_PAGES", 2)
        ocr = _FakeOcr()
        contents = _pdf(("text", "body"), ("image", ""), ("image", ""), ("image", ""))

        PdfParser(TextBudget(), ocr).parse(contents)

        assert ocr.read_pages == [1, 2]

    def test_without_ocr_a_scan_comes_back_with_no_text(self):
        pages = PdfParser(TextBudget()).parse(_pdf(("image", "invoice")))

        assert [p.text.strip() for p in pages] == [""]

    def test_when_ocr_is_not_installed_the_pages_stay_blank_instead_of_failing(self):
        pages = PdfParser(TextBudget(), _MissingOcr()).parse(_pdf(("image", "invoice")))

        assert [p.text.strip() for p in pages] == [""]

    @pytest.mark.skipif(shutil.which("tesseract") is None, reason="Tesseract is not installed")
    def test_real_tesseract_reads_a_scanned_page(self):
        pages = DocumentParser().parse(_pdf(("image", "Invoice total 4200")), "pdf")

        assert "4200" in pages[0].text


class TestExtractionTimeout:
    def test_pdfs_get_extra_time_for_ocr_but_other_files_do_not(self):
        timeout = ExtractionTimeout()

        assert timeout.seconds("pdf") == settings.EXTRACT_TIMEOUT_SECONDS + settings.OCR_EXTRA_TIMEOUT_SECONDS
        assert timeout.seconds("PDF") == timeout.seconds("pdf")
        assert timeout.seconds("docx") == settings.EXTRACT_TIMEOUT_SECONDS

    def test_no_extra_time_when_ocr_is_switched_off(self, monkeypatch):
        monkeypatch.setattr(settings, "OCR_ENABLED", False)

        assert ExtractionTimeout().seconds("pdf") == settings.EXTRACT_TIMEOUT_SECONDS
