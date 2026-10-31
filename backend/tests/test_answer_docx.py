import io

import docx

from src.utils.answer_docx import FOOTER_NOTE, AnswerDocxWriter

SOURCES = [{"n": 1, "name": "lease.pdf", "page": 3, "quote": "Rent is  $900\ndue monthly."}]


def _read(question: str | None, answer: str, sources: list[dict] | None = None) -> docx.document.Document:
    return docx.Document(io.BytesIO(AnswerDocxWriter().write(question, answer, sources or [])))


def _paragraphs(document) -> list[tuple[str, str]]:
    return [(p.style.name, p.text) for p in document.paragraphs]


def test_the_question_is_the_title_and_the_answer_follows():
    document = _read("What is the rent?", "The rent is $900 [1].")

    assert _paragraphs(document)[:2] == [("Heading 1", "What is the rent?"), ("Normal", "The rent is $900 [1].")]


def test_a_document_without_a_question_starts_with_the_answer():
    assert _paragraphs(_read(None, "Just the answer."))[0] == ("Normal", "Just the answer.")


def test_sources_are_listed_with_their_page_and_a_tidied_quote():
    document = _read("q", "a", SOURCES)

    assert ("Heading 2", "Sources") in _paragraphs(document)
    assert ("List Number", "lease.pdf, page 3: “Rent is $900 due monthly.”") in _paragraphs(document)


def test_a_source_without_a_page_leaves_the_page_out():
    document = _read("q", "a", [{"name": "notes.md", "page": None, "quote": "Call Sam."}])

    assert ("List Number", "notes.md: “Call Sam.”") in _paragraphs(document)


def test_every_export_ends_with_the_reminder_to_check_sources():
    assert _paragraphs(_read("q", "a"))[-1] == ("Normal", FOOTER_NOTE)


def test_bold_and_italic_text_keep_their_formatting_without_the_markdown_marks():
    document = _read(None, "A **firm** term and an *optional* one.")
    runs = [(run.text, run.bold, run.italic) for run in document.paragraphs[0].runs]

    assert runs == [
        ("A ", None, None),
        ("firm", True, None),
        (" term and an ", None, None),
        ("optional", None, True),
        (" one.", None, None),
    ]


def test_bullets_and_numbered_items_use_list_styles():
    document = _read(None, "- first\n- second\n\n1. one\n2. two")

    assert _paragraphs(document)[:4] == [
        ("List Bullet", "first"),
        ("List Bullet", "second"),
        ("List Number", "one"),
        ("List Number", "two"),
    ]


def test_a_markdown_link_keeps_its_text_and_drops_the_address():
    assert _paragraphs(_read(None, "See [the lease](https://example.com/x)."))[0][1] == "See the lease."


def test_headings_in_an_answer_become_word_headings():
    assert _paragraphs(_read(None, "## Terms\nBody"))[:2] == [("Heading 3", "Terms"), ("Normal", "Body")]


def test_a_markdown_table_becomes_a_word_table_with_a_bold_header():
    document = _read(None, "| Item | Amount |\n|---|---|\n| Rent | $900 |")
    [table] = document.tables

    assert [[cell.text for cell in row.cells] for row in table.rows] == [["Item", "Amount"], ["Rent", "$900"]]
    assert all(run.bold for cell in table.rows[0].cells for run in cell.paragraphs[0].runs)


def test_a_code_block_is_kept_as_one_monospaced_paragraph():
    document = _read(None, "```\nline one\nline two\n```")

    assert document.paragraphs[0].text == "line one\nline two"
    assert document.paragraphs[0].runs[0].font.name == "Courier New"
