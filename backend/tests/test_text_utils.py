import io
import zipfile

import docx
import fitz
import pytest

from src.config import settings
from src.core.exceptions import BadRequestError
from src.utils.text_utils import chunk_text, clean_text, extract_text_from_bytes


def _docx_bytes(build) -> bytes:
    document = docx.Document()
    build(document)
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


class TestExtractDocx:
    def test_keeps_paragraphs_and_table_cells_in_document_order(self):
        def build(document):
            document.add_paragraph("Payment terms")
            table = document.add_table(rows=2, cols=2)
            table.cell(0, 0).text = "Due"
            table.cell(0, 1).text = "30 days"
            table.cell(1, 0).text = "Late fee"
            table.cell(1, 1).text = "2%"
            document.add_paragraph("Termination")

        text = extract_text_from_bytes(_docx_bytes(build), "docx")

        assert text.splitlines() == ["Payment terms", "Due | 30 days", "Late fee | 2%", "Termination"]

    def test_empty_document_has_no_readable_text(self):
        assert clean_text(extract_text_from_bytes(_docx_bytes(lambda d: None), "docx")) == ""


def _zip_bytes(members: dict[str, bytes]) -> bytes:
    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w", zipfile.ZIP_DEFLATED) as archive:
        for name, data in members.items():
            archive.writestr(name, data)
    return buffer.getvalue()


class TestUploadSafetyLimits:
    def test_docx_that_unpacks_far_beyond_its_size_is_rejected(self):
        # Zeros compress to almost nothing, so the file is tiny but declares a huge unpacked size.
        bomb = _zip_bytes({"word/document.xml": b"\0" * (settings.MAX_DOCX_UNZIPPED_MB * 1024 * 1024 + 1)})

        assert len(bomb) < 1024 * 1024
        with pytest.raises(BadRequestError) as exc:
            extract_text_from_bytes(bomb, "docx")
        assert "too large" in exc.value.message

    def test_docx_with_too_many_entries_is_rejected(self):
        many = _zip_bytes({f"part{i}.xml": b"x" for i in range(settings.MAX_DOCX_ENTRIES + 1)})

        with pytest.raises(BadRequestError) as exc:
            extract_text_from_bytes(many, "docx")
        assert "too large" in exc.value.message

    def test_pdf_with_too_many_pages_is_rejected(self, monkeypatch):
        monkeypatch.setattr(settings, "MAX_PDF_PAGES", 2)
        pdf = fitz.open()
        for _ in range(3):
            pdf.new_page()
        contents = pdf.tobytes()
        pdf.close()

        with pytest.raises(BadRequestError) as exc:
            extract_text_from_bytes(contents, "pdf")
        assert "too large" in exc.value.message

    def test_text_longer_than_the_cap_is_rejected(self, monkeypatch):
        monkeypatch.setattr(settings, "MAX_EXTRACTED_CHARS", 10)

        with pytest.raises(BadRequestError) as exc:
            extract_text_from_bytes(b"a" * 11, "txt")
        assert "too large" in exc.value.message

    def test_text_at_the_cap_is_accepted(self, monkeypatch):
        monkeypatch.setattr(settings, "MAX_EXTRACTED_CHARS", 10)

        assert extract_text_from_bytes(b"a" * 10, "txt") == "a" * 10


class TestExtractOtherTypes:
    def test_txt_is_decoded(self):
        assert extract_text_from_bytes(b"hello", "txt") == "hello"

    def test_unsupported_type_raises(self):
        with pytest.raises(BadRequestError):
            extract_text_from_bytes(b"x", "exe")

    @pytest.mark.parametrize("ext", ["pdf", "docx"])
    def test_empty_or_corrupt_file_gives_a_plain_message(self, ext):
        for contents in (b"", b"not a real document"):
            with pytest.raises(BadRequestError) as exc:
                extract_text_from_bytes(contents, ext)
            assert "couldn't be read" in exc.value.message


def test_chunk_text_collapses_whitespace_and_respects_size():
    chunks = chunk_text("a  b\n\nc d e f", chunk_size=5, overlap=2)

    assert chunks[0] == "a b c"
    assert all(len(chunk) <= 5 for chunk in chunks)
