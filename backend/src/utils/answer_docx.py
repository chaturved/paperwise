import io
import re
from collections.abc import Sequence

import docx
from docx.document import Document as DocxDocument
from docx.text.paragraph import Paragraph

FOOTER_NOTE = "Exported from Paperwise. Check sources before relying on an answer."

_INLINE = re.compile(r"(\*\*[^*\n]+\*\*|\*[^*\n]+\*|_[^_\n]+_|`[^`\n]+`|\[[^\]\n]+\]\([^)\n]+\))")
_LINK = re.compile(r"\[([^\]\n]+)\]\([^)\n]+\)")
_HEADING = re.compile(r"^(#{1,3})\s+(.*)$")
_BULLET = re.compile(r"^\s*[-*+]\s+(.*)$")
_NUMBERED = re.compile(r"^\s*\d+[.)]\s+(.*)$")
_TABLE_ROW = re.compile(r"^\s*\|.*\|\s*$")
_TABLE_RULE = re.compile(r"^\s*\|?[\s:|-]+\|[\s:|-]*$")


class AnswerDocxWriter:
    """Writes one answer, with the question it answered and the passages it cites, as a Word document."""

    def write(self, question: str | None, answer: str, sources: Sequence[dict]) -> bytes:
        document = docx.Document()
        if question:
            document.add_heading(question, level=1)
        self._markdown(document, answer)
        if sources:
            document.add_heading("Sources", level=2)
            for source in sources:
                document.add_paragraph(self._source_line(source), style="List Number")
        note = document.add_paragraph()
        note.add_run(FOOTER_NOTE).italic = True

        buffer = io.BytesIO()
        document.save(buffer)
        return buffer.getvalue()

    @staticmethod
    def _source_line(source: dict) -> str:
        page = f", page {source['page']}" if source.get("page") else ""
        quote = " ".join(str(source.get("quote", "")).split())
        return f"{source.get('name', 'Document')}{page}: “{quote}”"

    def _markdown(self, document: DocxDocument, text: str) -> None:
        lines = text.strip().splitlines()
        index = 0
        while index < len(lines):
            line = lines[index]
            if not line.strip():
                index += 1
            elif line.lstrip().startswith("```"):
                index = self._code_block(document, lines, index)
            elif _TABLE_ROW.match(line):
                index = self._table(document, lines, index)
            else:
                self._block(document, line)
                index += 1

    def _block(self, document: DocxDocument, line: str) -> None:
        heading = _HEADING.match(line)
        if heading:
            document.add_heading(self._plain(heading.group(2)), level=min(len(heading.group(1)) + 1, 3))
        elif bullet := _BULLET.match(line):
            self._inline(document.add_paragraph(style="List Bullet"), bullet.group(1))
        elif numbered := _NUMBERED.match(line):
            self._inline(document.add_paragraph(style="List Number"), numbered.group(1))
        else:
            self._inline(document.add_paragraph(), line.strip())

    @staticmethod
    def _code_block(document: DocxDocument, lines: list[str], start: int) -> int:
        end = start + 1
        body: list[str] = []
        while end < len(lines) and not lines[end].lstrip().startswith("```"):
            body.append(lines[end])
            end += 1
        run = document.add_paragraph().add_run("\n".join(body))
        run.font.name = "Courier New"
        return end + 1

    def _table(self, document: DocxDocument, lines: list[str], start: int) -> int:
        end = start
        rows: list[list[str]] = []
        while end < len(lines) and _TABLE_ROW.match(lines[end]):
            if not _TABLE_RULE.match(lines[end]):
                rows.append([cell.strip() for cell in lines[end].strip().strip("|").split("|")])
            end += 1
        if rows:
            width = max(len(row) for row in rows)
            table = document.add_table(rows=len(rows), cols=width, style="Table Grid")
            for row_index, row in enumerate(rows):
                for column, cell in enumerate(row):
                    paragraph = table.cell(row_index, column).paragraphs[0]
                    self._inline(paragraph, cell)
                    if row_index == 0:
                        for run in paragraph.runs:
                            run.bold = True
        return end

    def _inline(self, paragraph: Paragraph, text: str) -> None:
        for part in _INLINE.split(text):
            if not part:
                continue
            if part.startswith("**") and part.endswith("**"):
                paragraph.add_run(part[2:-2]).bold = True
            elif part.startswith("`") and part.endswith("`"):
                paragraph.add_run(part[1:-1]).font.name = "Courier New"
            elif part[0] in "*_" and part[-1] == part[0] and len(part) > 2:
                paragraph.add_run(part[1:-1]).italic = True
            else:
                paragraph.add_run(_LINK.sub(r"\1", part))

    @staticmethod
    def _plain(text: str) -> str:
        return re.sub(r"[*_`]", "", _LINK.sub(r"\1", text))
